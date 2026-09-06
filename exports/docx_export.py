from docx import Document
from resume.models import Resume


def export_docx(resume: Resume, output_path: str = "resume.docx") -> str:
    doc = Document()
    pi = resume.personal_info

    doc.add_heading(pi.full_name or "Your Name", level=0)
    if pi.headline:
        doc.add_paragraph(pi.headline)
    contact = " | ".join(filter(None, [pi.email, pi.phone, pi.location, pi.linkedin, pi.github, pi.portfolio]))
    if contact:
        doc.add_paragraph(contact)

    if resume.professional_summary:
        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph(resume.professional_summary)

    if resume.work_experience:
        doc.add_heading("Work Experience", level=1)
        for exp in resume.work_experience:
            doc.add_paragraph(f"{exp.job_title} — {exp.company} ({exp.start_date} - {exp.end_date})", style="Heading 2")
            for b in exp.bullets:
                doc.add_paragraph(b, style="List Bullet")

    if resume.education:
        doc.add_heading("Education", level=1)
        for edu in resume.education:
            doc.add_paragraph(f"{edu.degree} — {edu.institution} ({edu.start_date} - {edu.end_date})")

    if resume.technical_skills:
        doc.add_heading("Technical Skills", level=1)
        doc.add_paragraph(", ".join(resume.technical_skills))

    if resume.projects:
        doc.add_heading("Projects", level=1)
        for p in resume.projects:
            doc.add_paragraph(f"{p.title}: {p.description}")

    if resume.certifications:
        doc.add_heading("Certifications", level=1)
        for c in resume.certifications:
            doc.add_paragraph(f"{c.name} — {c.issuer} ({c.date})")

    doc.save(output_path)
    return output_path
